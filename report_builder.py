import json, re
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

STRICT_RULES={'font':'Times New Roman','font_size_pt':14,'line_spacing':1.5,'first_line_cm':1.25,'left_cm':3.0,'right_cm':1.5,'top_cm':2.0,'bottom_cm':2.0}

def extract_document_text(path):
    p=Path(path); s=p.suffix.lower()
    if s=='.pdf': return '\n'.join(page.extract_text() or '' for page in PdfReader(str(p)).pages)
    if s=='.docx':
        d=Document(str(p)); out=[x.text for x in d.paragraphs if x.text.strip()]
        for t in d.tables: out.extend(' | '.join(c.text.strip() for c in row.cells) for row in t.rows)
        return '\n'.join(out)
    return None

def parse_spec(raw):
    if isinstance(raw,dict): d=raw
    else:
        x=str(raw or '').strip(); x=re.sub(r'^```(?:json)?\s*','',x,flags=re.I); x=re.sub(r'\s*```$','',x)
        m=re.search(r'\{[\s\S]*\}',x)
        if not m: raise ValueError('Модель не вернула JSON спецификацию отчета')
        d=json.loads(m.group(0))
    if not isinstance(d,dict): raise ValueError('Спецификация отчета должна быть объектом')
    defaults={'title':'ОТЧЕТ','discipline':'','lab_number':'','topic':'','student':'','group':'9СК-21','city':'Санкт-Петербург','year':'2026','sections':[]}
    for k,v in defaults.items(): d.setdefault(k,v)
    return d

def cap(n,text): return ('Рисунок %d - %s'%(n,re.sub(r'^Рис\.?\s*','',str(text or '').strip(),flags=re.I))).rstrip('. ')

def fmt(p,first=True,align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.alignment=align; f=p.paragraph_format; f.left_indent=Cm(0); f.right_indent=Cm(0); f.first_line_indent=Cm(1.25 if first else 0); f.space_before=Pt(0); f.space_after=Pt(0); f.line_spacing=1.5

def font(run,bold=None,size=14):
    run.font.name='Times New Roman'; run.font.size=Pt(size)
    if bold is not None: run.bold=bold

def build_docx(d,out):
    doc=Document(); s=doc.sections[0]; s.top_margin=Cm(2); s.bottom_margin=Cm(2); s.left_margin=Cm(3); s.right_margin=Cm(1.5)
    for x in ['ПРАВИТЕЛЬСТВО САНКТ-ПЕТЕРБУРГА','КОМИТЕТ ПО НАУКЕ И ВЫСШЕЙ ШКОЛЕ','Санкт-Петербургское государственное бюджетное профессиональное образовательное учреждение','«Академия инженерных технологий и управления»','','Отделение «Информационные технологии»','',d['title'],'',f"по дисциплине «{d['discipline']}»",f"по лабораторному занятию № {d['lab_number']}",f"по теме «{d['topic']}»"]:
        p=doc.add_paragraph(); fmt(p,False,WD_ALIGN_PARAGRAPH.CENTER); font(p.add_run(x),x==d['title'])
    for _ in range(5): doc.add_paragraph()
    for x in ['Работу выполнил(а):',f"Студент(ка) группы {d['group']}",d['student'] or '________________________','(ФИО студента)']:
        p=doc.add_paragraph(); fmt(p,False,WD_ALIGN_PARAGRAPH.LEFT); font(p.add_run(x))
    for _ in range(6): doc.add_paragraph()
    for x in [d['city'],str(d['year'])]: p=doc.add_paragraph(); fmt(p,False,WD_ALIGN_PARAGRAPH.CENTER); font(p.add_run(x))
    doc.add_section(WD_SECTION.NEW_PAGE); n=0
    for sec in d['sections']:
        h=str(sec.get('heading') or sec.get('title') or '').strip()
        if h: p=doc.add_paragraph(); fmt(p); font(p.add_run(h.rstrip('. ')),True); doc.add_paragraph()
        for x in sec.get('paragraphs',[]) or []: p=doc.add_paragraph(); fmt(p); font(p.add_run(str(x)))
        for x in sec.get('steps',[]) or []: p=doc.add_paragraph(); fmt(p); font(p.add_run(str(x)))
        for td in sec.get('tables',[]) or []:
            rows=td if isinstance(td,list) else td.get('rows',[])
            if rows:
                t=doc.add_table(rows=len(rows),cols=max(len(r) for r in rows)); t.style='Table Grid'
                for ri,row in enumerate(rows):
                    for ci,v in enumerate(row): t.cell(ri,ci).text=str(v)
        for fig in sec.get('figures',[]) or []:
            n+=1; p=doc.add_paragraph(); fmt(p); font(p.add_run(fig.get('reference') or f'Результат измерения представлен на рисунке {n}.'))
            p=doc.add_paragraph(); fmt(p,False,WD_ALIGN_PARAGRAPH.CENTER); font(p.add_run(f'[ МЕСТО ДЛЯ РИСУНКА {n} ]'),True)
            p=doc.add_paragraph(); fmt(p,False,WD_ALIGN_PARAGRAPH.CENTER); font(p.add_run(cap(n,fig.get('caption')))); doc.add_paragraph()
        for x in sec.get('answers',[]) or []: p=doc.add_paragraph(); fmt(p); font(p.add_run(str(x)))
    for sec in doc.sections:
        sec.different_first_page_header_footer=True; p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.append(fld); font(r)
    doc.save(str(out))

def build_pdf(d,out):
    fontname='Helvetica'
    for fp in ['/System/Library/Fonts/Supplemental/Times New Roman.ttf','/Library/Fonts/Times New Roman.ttf']:
        if Path(fp).exists(): pdfmetrics.registerFont(TTFont('TimesNewRoman',fp)); fontname='TimesNewRoman'; break
    body=ParagraphStyle('body',fontName=fontname,fontSize=14,leading=21,alignment=TA_JUSTIFY,firstLineIndent=1.25*cm)
    center=ParagraphStyle('center',fontName=fontname,fontSize=14,leading=21,alignment=TA_CENTER)
    doc=SimpleDocTemplate(str(out),pagesize=A4,leftMargin=3*cm,rightMargin=1.5*cm,topMargin=2*cm,bottomMargin=2*cm); story=[]
    for x in ['ПРАВИТЕЛЬСТВО САНКТ-ПЕТЕРБУРГА','КОМИТЕТ ПО НАУКЕ И ВЫСШЕЙ ШКОЛЕ','Академия инженерных технологий и управления','',d['title'],'',f"по дисциплине «{d['discipline']}»",f"по лабораторному занятию № {d['lab_number']}",f"по теме «{d['topic']}»"]: story += [Paragraph(str(x),center),Spacer(1,5)]
    story += [Spacer(1,100)]
    for x in ['Работу выполнил(а):',f"Студент(ка) группы {d['group']}",d['student'] or '________________________','(ФИО студента)']: story += [Paragraph(str(x),body),Spacer(1,3)]
    story += [Spacer(1,140),Paragraph(d['city'],center),Paragraph(str(d['year']),center),PageBreak()]; n=0
    for sec in d['sections']:
        h=str(sec.get('heading') or sec.get('title') or '').strip()
        if h: story += [Paragraph(h.rstrip('. '),center),Spacer(1,14)]
        for x in sec.get('paragraphs',[]) or []: story.append(Paragraph(escape(str(x)),body))
        for x in sec.get('steps',[]) or []: story.append(Paragraph(escape(str(x)),body))
        for fig in sec.get('figures',[]) or []:
            n+=1; story += [Paragraph(escape(str(fig.get('reference') or f'Результат измерения представлен на рисунке {n}.')),body),Spacer(1,20),Paragraph(f'[ МЕСТО ДЛЯ РИСУНКА {n} ]',center),Spacer(1,20),Paragraph(escape(cap(n,fig.get('caption'))),center),Spacer(1,12)]
        for x in sec.get('answers',[]) or []: story.append(Paragraph(escape(str(x)),body))
    def footer(c,doc): c.saveState(); c.setFont(fontname,14); c.drawCentredString(A4[0]/2,1*cm,str(doc.page) if doc.page>1 else ''); c.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)

def build_report_files(raw,out_dir,base='Отчет'):
    d=parse_spec(raw); root=Path(out_dir); root.mkdir(parents=True,exist_ok=True); safe=re.sub(r'[^A-Za-zА-Яа-я0-9._-]+','_',base).strip('_') or 'Отчет'
    a=root/(safe+'.docx'); b=root/(safe+'.pdf'); build_docx(d,a); build_pdf(d,b); return [a,b]

def strict_system_prompt():
    return json.dumps({'rules':STRICT_RULES,'schema':{'title':'ОТЧЕТ','discipline':'string','lab_number':'string','topic':'string','student':'string','group':'string','city':'string','year':'string','sections':[{'heading':'string','paragraphs':['string'],'steps':['string'],'tables':[[['cell']]],'figures':[{'reference':'string','caption':'string'}],'answers':['string']}]},'constraints':['Return ONLY valid JSON.','Use only supplied sources and user conditions.','Never invent measurements or screenshots.','For unavailable screenshots create a placeholder and caption.','Preserve source order and terminology.']},ensure_ascii=False)